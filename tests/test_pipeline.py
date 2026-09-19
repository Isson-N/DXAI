import csv
import shutil
import zipfile
from pathlib import Path

import numpy as np
import pytest
from openpyxl import load_workbook

from dxaqc.cli import main
from dxaqc.io import read_image
from dxaqc.labels import REGION_HIP, REGION_SPINE
from dxaqc.model import StubModel
from dxaqc.pipeline import run
from dxaqc.report import COLUMNS

from .conftest import make_dicom

ROOT = Path(__file__).resolve().parent.parent


def by_name(rows):
    return {Path(r["path_to_study"]).name: r for r in rows}


def test_pipeline_handles_good_and_bad_files(study_dir):
    rows, errors = run(study_dir, StubModel())
    r = by_name(rows)
    assert set(r) == {"spine.dcm", "hip.dcm", "spine_copy.dcm", "inverted.dcm", "broken.dcm"}
    assert r["spine.dcm"]["anatomical_region"] == REGION_SPINE
    assert r["hip.dcm"]["anatomical_region"] == REGION_HIP
    assert r["spine.dcm"]["processing_status"] == "Success"
    assert r["spine.dcm"]["violation_type"] == "" and r["spine.dcm"]["quality_class"] == 0
    assert r["broken.dcm"]["processing_status"] == "Failure"
    assert r["spine.dcm"]["study_uid"] == r["hip.dcm"]["study_uid"] != ""
    messages = {Path(e["path"]).name: e["message"] for e in errors}
    assert "не DICOM" in messages["notes.txt"]
    assert "нет изображения" in messages["DICOMDIR.dcm"]
    assert "broken.dcm" in messages


def test_output_order_is_deterministic(study_dir):
    first, _ = run(study_dir, StubModel())
    second, _ = run(study_dir, StubModel())
    strip = lambda rows: [{k: v for k, v in r.items() if k != "time_of_processing"} for r in rows]
    assert strip(first) == strip(second)


def test_monochrome1_is_inverted(tmp_path):
    arr = np.array([[0, 100], [200, 50]], dtype=np.uint8)  # максимум содержимого ≠ 255
    p1 = make_dicom(tmp_path / "m2.dcm", 2, 2, pixels=arr)
    p2 = make_dicom(tmp_path / "m1.dcm", 2, 2, photometric="MONOCHROME1", pixels=arr)
    a = read_image(p1, tmp_path).pixels
    b = read_image(p2, tmp_path).pixels
    assert np.array_equal(a, 255 - b)


def test_zip_input_with_cyrillic_names(study_dir, tmp_path):
    archive = tmp_path / "Для теста.zip"
    with zipfile.ZipFile(archive, "w") as z:
        for p in study_dir.rglob("*"):
            if p.is_file():
                z.write(p, Path("Исследования") / p.relative_to(study_dir))
    rows, _ = run(archive, StubModel())
    assert len(rows) == 5
    assert all(r["path_to_study"].startswith("Исследования/") for r in rows)


@pytest.mark.parametrize("suffix", [".csv", ".xlsx"])
def test_cli_writes_organizer_format(study_dir, tmp_path, suffix):
    out = tmp_path / "out" / f"results{suffix}"
    assert main(["predict", "--stub", "--input", str(study_dir), "--output", str(out)]) == 0
    if suffix == ".csv":
        with out.open(encoding="utf-8") as f:
            header = next(csv.reader(f))
    else:
        header = [c.value for c in next(load_workbook(out).active.iter_rows(max_row=1))]
    assert header == COLUMNS
    assert (out.parent / "errors.csv").exists()


def test_cli_missing_input(tmp_path):
    assert main(["predict", "--stub", "--input", str(tmp_path / "nope"), "--output", str(tmp_path / "r.csv")]) == 2


@pytest.mark.skipif(not (ROOT / "data" / "test").exists(), reason="нет данных организатора")
def test_organizer_test_files(tmp_path):
    src = ROOT / "data" / "test"
    rows, errors = run(src, StubModel())
    regions = {Path(r["path_to_study"]).name: r["anatomical_region"] for r in rows}
    assert all(r["processing_status"] == "Success" for r in rows)
    assert regions["CR000000_ПОП.dcm"] == REGION_SPINE
    assert regions["CR000000_ППОБ.dcm"] == REGION_HIP
    assert regions["CR000001_ЛПОБ.dcm"] == REGION_HIP
